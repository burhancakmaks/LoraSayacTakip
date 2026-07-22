"""Import requested MASKI workbooks and link rows only by unique meter number."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import sqlite3
import unicodedata
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATABASE = ROOT / "data" / "binalar.db"
DOWNLOADS = Path.home() / "Downloads"
REQUESTED = [
    "37-50 ADA E BLOK 72 ADET MASKİ ABONELİK (1).XLS",
    "ŞİRE PAZARI MASKİ ABONELİKLERİ.xlsx",
    "53 ADA MASKİ ABONELERİ.xlsx",
    "49 ADA 301 ADET  MASKİ ABONELİK (1).xlsx",
    "41-134   341 ADET  maski abonelik.xlsx",
    "46 ADA KONUT MASKİ ABONELERİ.xlsx",
    "4.ETAP TS SAYAÇ NO.xlsx",
    "37-50 ADA A-B BLOK 344 ADETMASKİ ABONELİK.xlsx",
]


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", " ", str(value).strip())


def folded(value: Any) -> str:
    value = unicodedata.normalize("NFKD", clean(value)).upper()
    return "".join(char for char in value if not unicodedata.combining(char))


def meter_key(value: Any) -> str:
    raw = folded(value)
    if not raw or raw in {"-", "YOK", "BOS"} or "TAKIL" in raw or "SAYAC" in raw:
        return ""
    key = "".join(char for char in raw if char.isalnum())
    return key if len(key) >= 5 and any(char.isdigit() for char in key) else ""


def header_kind(value: Any) -> str:
    key = re.sub(r"[^A-Z0-9]+", " ", folded(value)).strip()
    if "SAYAC" in key and ("NO" in key or "NUMARA" in key or key == "SAYACNO"):
        return "sayac_no"
    if "ABONE" in key and ("NO" in key or "NOSU" in key):
        return "abone_no"
    if "SICIL" in key:
        return "sicil_no"
    if key in {"ADA", "ADA NO"}:
        return "ada"
    if "BLOK" in key:
        return "blok"
    if "BAGIMSIZ BOLUM" in key or "KAPI NO" in key or key == "DAIRE NO":
        return "daire"
    if "KAT" in key:
        return "kat"
    if "ODA" in key:
        return "oda_sayisi"
    if "NITELIG" in key or "KULLANILIS" in key:
        return "kullanim"
    if "AD SOYAD" in key or key in {"ADI SOYADI", "ABONE ADI"}:
        return "ad_soyad"
    return ""


def infer_ada(file_name: str, sheet_name: str) -> str:
    match = re.search(r"(\d+(?:-\d+)?)\s*ADA", folded(file_name))
    if match:
        return match.group(1)
    match = re.search(r"ADA[- ]?(\d+)", folded(sheet_name))
    return match.group(1) if match else ""


def infer_mahalle(address: str) -> str:
    match = re.search(r"(.+?)\s+MAH(?:ALLESI|ALESI|\.)", folded(address))
    return clean(match.group(1)) if match else ""


def make_record(path: Path, sheet: str, row_number: int, source: dict[str, Any]) -> dict[str, Any]:
    meter = clean(source.get("sayac_no"))
    identity = f"{path.name}|{sheet}|{row_number}|{meter}|{clean(source.get('blok'))}"
    return {
        "kayit_id": hashlib.sha256(identity.encode("utf-8")).hexdigest(),
        "excel_satir_no": row_number,
        "ada": clean(source.get("ada")),
        "blok": clean(source.get("blok")),
        "kat": clean(source.get("kat")),
        "daire": clean(source.get("daire")),
        "ad_soyad": clean(source.get("ad_soyad")),
        "abone_no": clean(source.get("abone_no")),
        "sayac_no": meter,
        "adres": clean(source.get("adres")),
        "mahalle": clean(source.get("mahalle")),
        "kaynak_dosya": path.name,
        "durum": clean(source.get("kullanim")),
        "sicil_no": clean(source.get("sicil_no")),
        "oda_sayisi": clean(source.get("oda_sayisi")),
        "kaynak": {"sheet": sheet, "row": row_number, **source},
    }


def parse_standard(path: Path, sheet) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    mapping: dict[str, int] = {}
    address = ""
    fallback_ada = infer_ada(path.name, sheet.title)
    for row_number, cells in enumerate(sheet.iter_rows(values_only=True), start=1):
        values = [clean(value) for value in cells]
        nonempty = [value for value in values if value]
        joined = " ".join(nonempty)
        if len(joined) >= 15 and ("MAHAL" in folded(joined) or "SOKAK" in folded(joined)):
            address = joined
        kinds = {header_kind(value): index for index, value in enumerate(values) if header_kind(value)}
        if "sayac_no" in kinds and len(kinds) >= 2:
            mapping = kinds
            continue
        if not mapping:
            continue
        source = {kind: values[index] if index < len(values) else "" for kind, index in mapping.items()}
        if not meter_key(source.get("sayac_no")):
            continue
        source.setdefault("ada", fallback_ada)
        source["adres"] = address
        source["mahalle"] = infer_mahalle(address)
        source["ham_satir"] = values
        records.append(make_record(path, sheet.title, row_number, source))
    return records


def parse_wide(path: Path, sheet) -> list[dict[str, Any]]:
    rows = [[clean(value) for value in row] for row in sheet.iter_rows(values_only=True)]
    if len(rows) < 4:
        return []
    header_index = next(
        (index for index, row in enumerate(rows) if sum(header_kind(value) == "sayac_no" for value in row) >= 2),
        None,
    )
    if header_index is None:
        return []
    records: list[dict[str, Any]] = []
    header = rows[header_index]
    block_row = rows[header_index - 1] if header_index else []
    for meter_column, value in enumerate(header):
        if header_kind(value) != "sayac_no":
            continue
        unit_column = meter_column - 1
        block = block_row[unit_column] if unit_column < len(block_row) else ""
        for row_index in range(header_index + 1, len(rows)):
            meter = rows[row_index][meter_column] if meter_column < len(rows[row_index]) else ""
            if not meter_key(meter):
                continue
            unit = rows[row_index][unit_column] if unit_column < len(rows[row_index]) else ""
            source = {
                "ada": infer_ada(path.name, sheet.title),
                "blok": block,
                "daire": unit,
                "sayac_no": meter,
                "kullanim": "ORTAK ALAN" if "2''" in unit or folded(unit) == "KD" else "DAİRE",
                "ham_satir": rows[row_index],
            }
            records.append(make_record(path, sheet.title, row_index + 1, source))
    return records


def parse_workbook(path: Path) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    with path.open("rb") as stream:
        workbook = load_workbook(stream, read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            header_rows = 0
            for row in sheet.iter_rows(values_only=True):
                header_rows = max(header_rows, sum(header_kind(value) == "sayac_no" for value in row))
            parser = parse_wide if header_rows >= 2 else parse_standard
            records.extend(parser(path, sheet))
    return records


def ensure_tables(db: sqlite3.Connection) -> None:
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS excel_import_batches (
          id INTEGER PRIMARY KEY AUTOINCREMENT, source_file TEXT NOT NULL,
          source_sha256 TEXT NOT NULL UNIQUE, generated_at TEXT NOT NULL,
          imported_at TEXT NOT NULL DEFAULT (datetime('now')),
          total_rows INTEGER NOT NULL, valid_rows INTEGER NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_excel_sayac_no ON excel_abonelikler(sayac_no);
        CREATE INDEX IF NOT EXISTS idx_excel_bina_id ON excel_abonelikler(bina_id);
        """
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true", help="Write the verified integration to SQLite")
    parser.add_argument("--verify", action="store_true", help="Verify the current SQLite integration")
    args = parser.parse_args()

    paths = [DOWNLOADS / name for name in REQUESTED]
    missing = [str(path) for path in paths if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing requested files: " + ", ".join(missing))

    by_file = {path: parse_workbook(path) for path in paths}
    all_records = [record for records in by_file.values() for record in records]

    with sqlite3.connect(DATABASE) as db:
        db.row_factory = sqlite3.Row
        meter_rows: dict[str, list[sqlite3.Row]] = defaultdict(list)
        for row in db.execute(
            """SELECT s.* FROM sayac s JOIN binalar b ON b.id=s.bina_id
            WHERE TRIM(COALESCE(s.sayac_id, '')) <> ''
              AND b.layer <> 'MASKI_EXCEL_ABONELIK_YAKLASIK'"""
        ):
            key = meter_key(row["sayac_id"])
            if key:
                meter_rows[key].append(row)

        matched = 0
        ambiguous = 0
        unmatched = 0
        matched_buildings: set[int] = set()
        for record in all_records:
            candidates = meter_rows.get(meter_key(record["sayac_no"]), [])
            building_ids = {int(row["bina_id"]) for row in candidates}
            if len(building_ids) == 1:
                record["bina_id"] = next(iter(building_ids))
                record["sayac_no"] = clean(candidates[0]["sayac_id"])
                record["sayac_rows"] = [int(row["id"]) for row in candidates if int(row["bina_id"]) == record["bina_id"]]
                matched += 1
                matched_buildings.add(record["bina_id"])
            elif building_ids:
                record["bina_id"] = None
                record["sayac_rows"] = []
                ambiguous += 1
            else:
                record["bina_id"] = None
                record["sayac_rows"] = []
                unmatched += 1

    summary = {
        "mode": "apply" if args.apply else "dry-run",
        "files": {path.name: len(records) for path, records in by_file.items()},
        "valid_meter_rows": len(all_records),
        "matched_rows": matched,
        "matched_buildings": len(matched_buildings),
        "ambiguous_rows": ambiguous,
        "unmatched_rows": unmatched,
    }
    if args.verify:
        placeholders = ",".join("?" for _ in REQUESTED)
        with sqlite3.connect(DATABASE) as db:
            linked, linked_buildings, broken_links, address_rows = db.execute(
                f"""
                SELECT COUNT(*), COUNT(DISTINCT e.bina_id),
                  SUM(CASE WHEN NOT EXISTS (
                    SELECT 1 FROM sayac s
                    WHERE s.bina_id=e.bina_id AND s.sayac_id=e.sayac_no AND s.sayac_id<>''
                  ) THEN 1 ELSE 0 END),
                  SUM(CASE WHEN TRIM(e.adres)<>'' THEN 1 ELSE 0 END)
                FROM excel_abonelikler e
                WHERE e.kaynak_dosya IN ({placeholders}) AND e.bina_id IS NOT NULL
                """,
                REQUESTED,
            ).fetchone()
        summary["database_verification"] = {
            "linked_rows": linked,
            "linked_buildings": linked_buildings,
            "links_without_exact_meter": broken_links,
            "linked_rows_with_address": address_rows,
            "green_rule": "has_meter_number",
        }
        print(json.dumps(summary, ensure_ascii=False, indent=2))
        return
    if not args.apply:
        print(json.dumps(summary, ensure_ascii=False, indent=2))
        return

    backup = DATABASE.with_name(f"binalar.before-meter-only-import-{datetime.now().strftime('%Y%m%d-%H%M%S')}.db")
    shutil.copy2(DATABASE, backup)
    with sqlite3.connect(DATABASE, timeout=30) as db:
        ensure_tables(db)
        db.execute("BEGIN IMMEDIATE")
        try:
            for path, records in by_file.items():
                digest = hashlib.sha256(path.read_bytes()).hexdigest()
                generated_at = datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat()
                db.execute(
                    """
                    INSERT INTO excel_import_batches(source_file, source_sha256, generated_at, total_rows, valid_rows)
                    VALUES (?, ?, ?, ?, ?)
                    ON CONFLICT(source_sha256) DO UPDATE SET source_file=excluded.source_file,
                      generated_at=excluded.generated_at, imported_at=datetime('now'),
                      total_rows=excluded.total_rows, valid_rows=excluded.valid_rows
                    """,
                    (path.name, digest, generated_at, len(records), len(records)),
                )
                batch_id = db.execute(
                    "SELECT id FROM excel_import_batches WHERE source_sha256=?", (digest,)
                ).fetchone()[0]
                db.execute("DELETE FROM excel_abonelikler WHERE batch_id=?", (batch_id,))
                db.executemany(
                    """
                    INSERT INTO excel_abonelikler(
                      kayit_id,batch_id,excel_satir_no,ada,blok,kat,daire,ad_soyad,
                      abone_no,sayac_no,adres,mahalle,kaynak_dosya,durum,kaynak_json,bina_id
                    ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                    ON CONFLICT(kayit_id) DO UPDATE SET batch_id=excluded.batch_id,
                      ada=excluded.ada,blok=excluded.blok,kat=excluded.kat,daire=excluded.daire,
                      ad_soyad=excluded.ad_soyad,abone_no=excluded.abone_no,
                      sayac_no=excluded.sayac_no,adres=excluded.adres,mahalle=excluded.mahalle,
                      kaynak_dosya=excluded.kaynak_dosya,durum=excluded.durum,
                      kaynak_json=excluded.kaynak_json,bina_id=excluded.bina_id,updated_at=datetime('now')
                    """,
                    [
                        (
                            r["kayit_id"], batch_id, r["excel_satir_no"], r["ada"], r["blok"],
                            r["kat"], r["daire"], r["ad_soyad"], r["abone_no"], r["sayac_no"],
                            r["adres"], r["mahalle"], r["kaynak_dosya"], r["durum"],
                            json.dumps(r["kaynak"], ensure_ascii=False), r["bina_id"],
                        )
                        for r in records
                    ],
                )

            for record in all_records:
                if record["bina_id"] is None:
                    continue
                for meter_row_id in record["sayac_rows"]:
                    db.execute(
                        """
                        UPDATE sayac SET
                          blok_no=CASE WHEN TRIM(COALESCE(blok_no,''))='' THEN ? ELSE blok_no END,
                          kat=CASE WHEN TRIM(COALESCE(kat,''))='' THEN ? ELSE kat END,
                          kapi_no=CASE WHEN TRIM(COALESCE(kapi_no,''))='' THEN ? ELSE kapi_no END,
                          daire_no=CASE WHEN TRIM(COALESCE(daire_no,''))='' THEN ? ELSE daire_no END,
                          oda_sayisi=CASE WHEN TRIM(COALESCE(oda_sayisi,'')) IN ('','YOK') AND ?<>'' THEN ? ELSE oda_sayisi END,
                          kullanilis_sekli=CASE WHEN TRIM(COALESCE(kullanilis_sekli,'')) IN ('','DAÝRE','DAÄ°RE') AND ?<>'' THEN ? ELSE kullanilis_sekli END,
                          sicil_no=CASE WHEN TRIM(COALESCE(sicil_no,''))='' THEN ? ELSE sicil_no END,
                          abone_no=CASE WHEN TRIM(COALESCE(abone_no,''))='' THEN ? ELSE abone_no END,
                          updated_at=datetime('now') WHERE id=?
                        """,
                        (
                            record["blok"], record["kat"], record["daire"], record["daire"],
                            record["oda_sayisi"], record["oda_sayisi"], record["durum"], record["durum"],
                            record["sicil_no"], record["abone_no"], meter_row_id,
                        ),
                    )
            db.commit()
        except Exception:
            db.rollback()
            raise

    summary["backup"] = str(backup)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
